#!/usr/bin/env python3
"""
Карта 301 со старого invit.by — из утверждённой таблицы клиента
(files/invit_старые_URL_заполнено.xlsx, колонки «Старый URL» и «Новый URL»)
в invit-theme/data/redirects.json. Тема читает её в inc/redirects.php.

exact  — адрес целиком (без хвостового «/», в нижнем регистре; у старых
         index.php?route=… — вместе с параметрами);
prefix — категории старого сайта: любой их неизвестный подадрес ведёт туда же.
"""
import json
import sys
from pathlib import Path
from urllib.parse import unquote, urlsplit

from openpyxl import load_workbook

SRC = Path(sys.argv[1] if len(sys.argv) > 1 else 'files/invit_старые_URL_заполнено.xlsx')
DST = Path('invit-theme/data/redirects.json')


def key(url):
    parts = urlsplit(url.strip())
    path = unquote(parts.path).strip().lower()
    if len(path) > 1:
        path = path.rstrip('/')
    query = unquote(parts.query).strip().lower()
    return path + ('?' + query if query else '')


def target(url):
    parts = urlsplit(url.strip())
    return parts.path + ('?' + parts.query if parts.query else '')


ws = load_workbook(SRC, read_only=True).active
header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
col_kind, col_old, col_new = header.index('Раздел'), header.index('Старый URL'), header.index('Новый URL (куда 301)')

exact, prefix = {}, {}
for row in ws.iter_rows(min_row=2, values_only=True):
    old, new = row[col_old], row[col_new]
    if not old or not new:
        continue
    exact[key(old)] = target(new)
    if 'категория' in (row[col_kind] or ''):
        prefix[key(old)] = target(new)

DST.write_text(json.dumps({'exact': exact, 'prefix': prefix}, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
print(f'{DST}: {len(exact)} адресов, {len(prefix)} старых категорий')
