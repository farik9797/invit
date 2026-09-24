"""Снимки товарам, у которых их нет, из архива поставщика.

Половина каталога пришла без фото: в момент импорта у строки не было кода
OKDP под рукой — коды лежат в выгрузках поставщика, а не в сведённом csv.
Здесь собираем артикул -> OKDP по всем выгрузкам и достаём снимки из
`files/by-full.zip` тем же способом, что и при импорте разделов.

Чужие сайты для этого не годятся: там снимки принадлежат их владельцам,
а у нас есть архив самого поставщика.

Запуск: python3 scripts/fill-photos-from-archive.py
Скрипт идемпотентен: строки, где фото уже есть, не трогает.
"""
import csv, re, shutil, sys, zipfile
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
WOO = ROOT / 'files/woocommerce_import_variations.csv'
ARCHIVE = ROOT / 'files/by-full.zip'
IMAGES = ROOT / 'files/images'
BACKUP = WOO.with_suffix('.csv.before-photos')

# Выгрузки поставщика, где есть пара «артикул — код OKDP»
SOURCES = [
    ROOT / 'files/export_old_2026-09-01.csv',
    ROOT / 'files/Болты_гайки_шайбы_2026-09-17.xlsx',
]


def plain(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def codes_from(path):
    """Пары «артикул -> OKDP» из csv или xlsx выгрузки."""
    if path.suffix.lower() == '.csv':
        csv.field_size_limit(10 ** 7)
        with path.open(encoding='utf-8-sig', errors='replace', newline='') as f:
            for row in csv.DictReader(f, delimiter=';'):
                yield plain(row.get('ARTIKUL')), plain(row.get('OKDP'))
        return

    from openpyxl import load_workbook
    sheet = load_workbook(path, read_only=True).active
    rows = sheet.iter_rows(values_only=True)
    header = [plain(h) for h in next(rows)]
    for line in rows:
        row = dict(zip(header, line))
        yield plain(row.get('ARTIKUL')), plain(row.get('OKDP'))


PACKAGING = re.compile(
    r'\((?=[^)]*(?:шт|уп\.|упак|короб|пакет|зип-лок|мешк|ведр))[^)]*\)', re.I)


def pack_free(name):
    """Название без фасовки: «Саморез 4.2х19 (50 шт в зип-локе)» -> «саморез 4.2х19»."""
    return ' '.join(PACKAGING.sub(' ', name).split()).lower()


def photos_for(archive, entries, sku):
    names = []
    for number, entry in enumerate(entries, 1):
        name = f'{sku}-{number}.jpg'
        target = IMAGES / name
        if not target.exists():
            with archive.open(entry) as src, target.open('wb') as dst:
                shutil.copyfileobj(src, dst)
        names.append(name)
    return names


def main():
    IMAGES.mkdir(exist_ok=True)

    codes = {}
    for path in SOURCES:
        if not path.exists():
            print(f'нет файла: {path.name}')
            continue
        for sku, code in codes_from(path):
            if sku and code:
                codes.setdefault(sku, code)
    print(f'артикулов с кодом OKDP: {len(codes)}')

    archive = zipfile.ZipFile(ARCHIVE)
    pictures = defaultdict(list)
    for name in archive.namelist():
        if '/' in name and name.lower().endswith(('.jpg', '.jpeg', '.png')):
            pictures[name.split('/')[0]].append(name)
    for entries in pictures.values():
        entries.sort()
    print(f'папок со снимками в архиве: {len(pictures)}')

    csv.field_size_limit(10 ** 7)
    with WOO.open(encoding='utf-8-sig', newline='') as f:
        rows = list(csv.DictReader(f))
    columns = list(rows[0].keys())

    empty = [r for r in rows if not r['Images'].strip()]
    filled = 0
    for row in empty:
        sku = row['SKU'].strip()
        entries = pictures.get(codes.get(sku, ''), [])
        if not entries:
            continue
        row['Images'] = ', '.join(photos_for(archive, entries, sku))
        filled += 1

    # Вторым заходом — та же позиция в другой фасовке: «(1000 шт в карт. уп.)»
    # и «(50 шт в зип-локе)» это один и тот же саморез, снимок у них общий.
    # Скобки убираем только там, где внутри действительно про упаковку: у
    # отрезного круга «(по металлу)» и «(по камню)» — разные товары.
    known = {}
    for row in rows:
        if row['Images'].strip():
            known.setdefault(pack_free(row['Name']), row['Images'])

    shared = 0
    for row in rows:
        if row['Images'].strip():
            continue
        images = known.get(pack_free(row['Name']))
        if images:
            row['Images'] = images
            shared += 1
    filled += shared

    if filled:
        if not BACKUP.exists():
            BACKUP.write_bytes(WOO.read_bytes())
        with WOO.open('w', encoding='utf-8-sig', newline='') as f:
            writer = csv.DictWriter(f, fieldnames=columns, lineterminator='\r\n')
            writer.writeheader()
            writer.writerows(rows)

    print(f'строк без фото : {len(empty)}')
    print(f'добавлено фото : {filled} (из них {shared} — общие с другой фасовкой)')
    print(f'осталось без   : {len(empty) - filled}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
