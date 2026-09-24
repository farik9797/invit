"""Метрический крепёж из выгрузки «Болты, гайки, шайбы» в сведённый импорт.

Выгрузка приходит в xlsx с теми же колонками, что и старая: разделы тремя
PARENTID_NAME0..2, название в NAIMEN, артикул в ARTIKUL, характеристики
в XARAKT. Раздел подбираем по названию: болт, гайка, шайба, шпилька.

Цену не переносим: в этой выгрузке только оптовая без НДС, а это другая
цена и покупателю она не адресована — в карточке будет «цена по запросу».

Фото берём из by-full.zip по коду OKDP, как в build-from-old.py.

Запуск: python3 scripts/build-metric-fasteners.py
Скрипт идемпотентен: известные артикулы пропускаются.
"""
import csv, html, re, shutil, sys, zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
HERE = ROOT / 'files'
SRC = HERE / 'Болты_гайки_шайбы_2026-09-17.xlsx'
WOO = HERE / 'woocommerce_import_variations.csv'
ARCHIVE = HERE / 'by-full.zip'
IMAGES = HERE / 'images'
BACKUP = WOO.with_suffix('.csv.before-metric')

# Порядок важен: шпилька проверяется раньше болта, иначе «шпилька резьбовая
# с болтом» ушла бы не туда.
TARGETS = [
    (re.compile(r'шпилька', re.I), 'Крепёж > Шпилька резьбовая'),
    (re.compile(r'\bгайк', re.I), 'Крепёж > Гайки'),
    (re.compile(r'\bшайб', re.I), 'Крепёж > Шайбы'),
    (re.compile(r'\bболт|\bвинт', re.I), 'Крепёж > Болты'),
]


def get(row, key):
    value = row.get(key)
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def as_html(text):
    lines = [line.strip() for line in str(text).splitlines()]
    return ''.join(f'<p>{html.escape(line)}</p>' for line in lines if line)


def num(value, factor=1, digits=2):
    raw = str(value).replace(',', '.').strip()
    try:
        value = round(float(raw) * factor, digits)
    except ValueError:
        return ''
    return f'{value:g}' if value else ''


def target_of(name):
    for pattern, path in TARGETS:
        if pattern.search(name):
            return path
    return ''


def photos_for(archive, code, sku):
    if not code:
        return []
    entries = sorted(n for n in archive.namelist()
                     if n.startswith(f'{code}/') and n.lower().endswith(('.jpg', '.jpeg', '.png')))
    names = []
    for number, entry in enumerate(entries, 1):
        name = f'{sku}-{number}.jpg'
        target = IMAGES / name
        if not target.exists():
            with archive.open(entry) as src, target.open('wb') as dst:
                shutil.copyfileobj(src, dst)
        names.append(name)
    return names


def to_woo(row, columns, archive, target):
    sku = get(row, 'ARTIKUL')
    values = {
        'Type': 'simple',
        'SKU': sku,
        'Name': get(row, 'NAIMEN'),
        'Published': '1',
        'Is featured?': '0',
        'Visibility in catalog': 'visible',
        'Short description': '',
        'Description': as_html(get(row, 'XARAKT')),
        'Tax status': 'taxable',
        'Tax class': '',
        'In stock?': '1',
        'Weight (kg)': num(get(row, 'VESB'), digits=3),
        'Length (cm)': num(get(row, 'DLINA'), factor=100),
        'Width (cm)': num(get(row, 'SHIRINA'), factor=100),
        'Height (cm)': num(get(row, 'VISOTA'), factor=100),
        'Regular price': '',
        'Categories': target,
        'Tags': '',
        'Images': ', '.join(photos_for(archive, get(row, 'OKDP'), sku)),
        'Attribute 1 name': 'Бренд',
        'Attribute 1 value(s)': get(row, 'BRAND'),
        'Attribute 1 visible': '1',
        'Attribute 1 global': '0',
        'Attribute 2 name': 'Страна',
        'Attribute 2 value(s)': get(row, 'STRANA'),
        'Attribute 2 visible': '1',
        'Attribute 2 global': '0',
    }
    return {column: values.get(column, '') for column in columns}


def supplier_rows():
    from openpyxl import load_workbook
    sheet = load_workbook(SRC, read_only=True).active
    stream = sheet.iter_rows(values_only=True)
    header = [str(h) for h in next(stream)]
    return [dict(zip(header, line)) for line in stream if line and line[5]]


def main():
    IMAGES.mkdir(exist_ok=True)
    supplier = supplier_rows()

    rows = list(csv.DictReader(WOO.open(encoding='utf-8-sig')))
    columns = list(rows[0].keys())
    known = {r['SKU'].strip() for r in rows}

    archive = zipfile.ZipFile(ARCHIVE)
    added, skipped, unrouted = [], 0, []
    for row in supplier:
        sku, name = get(row, 'ARTIKUL'), get(row, 'NAIMEN')
        if not sku or sku in known:
            skipped += 1
            continue
        target = target_of(name)
        if not target:
            unrouted.append(name)
            continue
        added.append(to_woo(row, columns, archive, target))
        known.add(sku)

    if added:
        if not BACKUP.exists():
            BACKUP.write_bytes(WOO.read_bytes())
        with WOO.open('w', encoding='utf-8-sig', newline='') as f:
            writer = csv.DictWriter(f, fieldnames=columns, lineterminator='\r\n')
            writer.writeheader()
            writer.writerows(rows + added)

    print(f'в выгрузке : {len(supplier)}')
    print(f'добавлено  : {len(added)} (уже были: {skipped})')
    from collections import Counter
    for path, count in Counter(r['Categories'] for r in added).most_common():
        print(f'   {count:5}  {path}')
    print(f'со снимком : {sum(1 for r in added if r["Images"])}')
    if unrouted:
        print(f'не разложено: {len(unrouted)}')
        for name in unrouted[:10]:
            print('   -', name[:70])
    print(f'→ {WOO.relative_to(ROOT)}: {len(rows)} + {len(added)} = {len(rows) + len(added)}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
